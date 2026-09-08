"""Local search & file analysis (Constitution sections 5A and 5D).

Everything works offline. Optional libraries (pypdf, python-docx, openpyxl,
Pillow) unlock richer formats when installed; without them Vidit still
handles text, code, markdown, JSON, CSV and can list what is in archives.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from ..guardian import Capability
from ..utils import human_size, keywords, truncate
from .base import Tool, ToolContext, ToolResult

log = logging.getLogger("vidit.tools.files")

TEXT_EXT = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".xml", ".html", ".htm", ".css", ".js", ".ts", ".tsx", ".jsx", ".py", ".java", ".c", ".h", ".cpp", ".hpp",
    ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql", ".kt", ".swift", ".dart", ".lua", ".r",
    ".ipynb", ".env", ".gitignore", ".dockerfile",
}
CODE_EXT = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".rb",
            ".php", ".sh", ".bat", ".ps1", ".sql", ".kt", ".swift", ".dart", ".lua", ".r"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".aac"}
VIDEO_EXT = {".mp4", ".mkv", ".mov", ".avi", ".webm"}
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", ".cache", "AppData", "$Recycle.Bin",
             "Windows", "Program Files", "Program Files (x86)", ".idea", ".vscode", "dist", "build", "target"}

MAX_READ_BYTES = 2_000_000


# ---------------------------------------------------------------------------
# reading
# ---------------------------------------------------------------------------
def read_text_file(path: Path, limit: int = MAX_READ_BYTES) -> str:
    raw = path.read_bytes()[:limit]
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path: Path) -> Dict[str, Any]:
    """Return {"text": str, "kind": str, "meta": {...}} for many file types."""
    suffix = path.suffix.lower()
    meta: Dict[str, Any] = {"name": path.name, "size": human_size(path.stat().st_size), "suffix": suffix}

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(str(path))
            pages = [(p.extract_text() or "") for p in reader.pages[:60]]
            meta["pages"] = len(reader.pages)
            return {"text": "\n\n".join(pages), "kind": "document", "meta": meta}
        except ImportError:
            return {"text": "", "kind": "document", "meta": meta,
                    "note": "Install `pypdf` so I can read PDFs (pip install pypdf)."}
        except Exception as exc:  # noqa: BLE001
            return {"text": "", "kind": "document", "meta": meta, "note": f"PDF unreadable: {exc}"}

    if suffix == ".docx":
        try:
            import docx  # type: ignore

            document = docx.Document(str(path))
            text = "\n".join(p.text for p in document.paragraphs)
            return {"text": text, "kind": "document", "meta": meta}
        except ImportError:
            # docx is a zip of XML; fall back to a crude extraction.
            try:
                with zipfile.ZipFile(path) as zf:
                    xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
                text = re.sub(r"<[^>]+>", " ", xml)
                return {"text": " ".join(text.split()), "kind": "document", "meta": meta}
            except Exception as exc:  # noqa: BLE001
                return {"text": "", "kind": "document", "meta": meta, "note": f"DOCX unreadable: {exc}"}

    if suffix in (".xlsx", ".xlsm"):
        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:5]:
                out.append(f"## Sheet: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 200:
                        out.append("…")
                        break
                    out.append(", ".join("" if v is None else str(v) for v in row))
            return {"text": "\n".join(out), "kind": "spreadsheet", "meta": meta}
        except ImportError:
            return {"text": "", "kind": "spreadsheet", "meta": meta, "note": "Install `openpyxl` so I can read Excel files."}

    if suffix in (".csv", ".tsv"):
        text = read_text_file(path)
        dialect = "excel-tab" if suffix == ".tsv" else "excel"
        rows = list(csv.reader(io.StringIO(text), dialect=dialect))
        meta["rows"] = len(rows)
        meta["columns"] = rows[0] if rows else []
        stats = _column_stats(rows)
        return {"text": text[:20000], "kind": "data", "meta": meta, "stats": stats}

    if suffix == ".json":
        text = read_text_file(path)
        try:
            data = json.loads(text)
            meta["json_type"] = type(data).__name__
            if isinstance(data, dict):
                meta["keys"] = list(data)[:30]
            elif isinstance(data, list):
                meta["items"] = len(data)
        except json.JSONDecodeError:
            pass
        return {"text": text[:20000], "kind": "data", "meta": meta}

    if suffix == ".ipynb":
        try:
            nb = json.loads(read_text_file(path))
            cells = []
            for cell in nb.get("cells", []):
                src = "".join(cell.get("source", []))
                cells.append(f"# [{cell.get('cell_type')}]\n{src}")
            return {"text": "\n\n".join(cells), "kind": "code", "meta": meta}
        except json.JSONDecodeError:
            pass

    if suffix in (".zip",):
        try:
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
            meta["files"] = len(names)
            return {"text": "\n".join(names[:300]), "kind": "archive", "meta": meta}
        except zipfile.BadZipFile as exc:
            return {"text": "", "kind": "archive", "meta": meta, "note": str(exc)}

    if suffix in IMAGE_EXT:
        try:
            from PIL import Image  # type: ignore

            with Image.open(path) as img:
                meta.update({"width": img.width, "height": img.height, "mode": img.mode, "format": img.format})
        except ImportError:
            meta["note"] = "Install `Pillow` for image details."
        except Exception as exc:  # noqa: BLE001
            meta["note"] = f"image unreadable: {exc}"
        return {"text": "", "kind": "image", "meta": meta}

    if suffix in AUDIO_EXT:
        return {"text": "", "kind": "audio", "meta": meta, "note": "Transcribe with the ears module (faster-whisper)."}
    if suffix in VIDEO_EXT:
        return {"text": "", "kind": "video", "meta": meta, "note": "Frame extraction needs OpenCV (opencv-python)."}

    if suffix in TEXT_EXT or suffix == "" or path.stat().st_size < 200_000:
        kind = "code" if suffix in CODE_EXT else "text"
        return {"text": read_text_file(path), "kind": kind, "meta": meta}

    return {"text": "", "kind": "binary", "meta": meta}


def _column_stats(rows: List[List[str]]) -> Dict[str, Any]:
    if len(rows) < 2:
        return {}
    header, body = rows[0], rows[1:]
    stats: Dict[str, Any] = {}
    for idx, col in enumerate(header[:20]):
        values = [r[idx] for r in body if idx < len(r) and r[idx] != ""]
        numeric: List[float] = []
        for v in values:
            try:
                numeric.append(float(v.replace(",", "")))
            except ValueError:
                break
        if numeric and len(numeric) == len(values):
            stats[col] = {"type": "number", "min": min(numeric), "max": max(numeric),
                          "mean": round(sum(numeric) / len(numeric), 3), "count": len(numeric)}
        else:
            stats[col] = {"type": "text", "unique": len(set(values)), "count": len(values)}
    return stats


# ---------------------------------------------------------------------------
# local search
# ---------------------------------------------------------------------------
class LocalSearch:
    """Search file names and contents inside the allowed folders."""

    def __init__(self, roots_provider, is_private) -> None:
        self._roots = roots_provider
        self._is_private = is_private

    def iter_files(self, roots: Optional[Iterable[Path]] = None, max_files: int = 20000):
        count = 0
        for root in roots or self._roots():
            root = Path(root).expanduser()
            if not root.exists():
                continue
            for dirpath, dirnames, filenames in os.walk(root):
                dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
                if self._is_private(dirpath):
                    dirnames[:] = []
                    continue
                for name in filenames:
                    count += 1
                    if count > max_files:
                        return
                    yield Path(dirpath) / name

    def search(self, query: str, limit: int = 20, content: bool = True) -> List[Dict[str, Any]]:
        terms = keywords(query) or [query.lower()]
        results: List[Dict[str, Any]] = []
        for path in self.iter_files():
            score = 0.0
            lname = path.name.lower()
            for term in terms:
                if term in lname:
                    score += 2
            snippet = ""
            if content and path.suffix.lower() in TEXT_EXT and path.stat().st_size < 1_500_000:
                try:
                    text = read_text_file(path, 300_000).lower()
                except OSError:
                    text = ""
                hits = sum(text.count(t) for t in terms)
                if hits:
                    score += min(5, hits * 0.5)
                    pos = min((text.find(t) for t in terms if t in text), default=-1)
                    if pos >= 0:
                        snippet = truncate(text[max(0, pos - 80): pos + 160], 240)
            if score > 0:
                results.append({"path": str(path), "score": score, "snippet": snippet, "size": human_size(path.stat().st_size)})
        results.sort(key=lambda r: r["score"], reverse=True)
        return results[:limit]


# ---------------------------------------------------------------------------
# analyzer
# ---------------------------------------------------------------------------
class FileAnalyzer:
    def __init__(self, llm) -> None:
        self.llm = llm

    def analyze(self, path: Path, question: str = "") -> Dict[str, Any]:
        info = extract_text(path)
        text = info.get("text", "")
        summary = ""
        if text.strip():
            prompt = (f"File: {path.name}\nKind: {info['kind']}\n\nContent (may be truncated):\n{text[:12000]}\n\n"
                      + (f"Question: {question}" if question else "Summarise the key points, then list anything notable."))
            summary = self.llm.quick(prompt, "You are Vidit analysing a file for your brother. Be concrete and useful.")
        elif info["kind"] == "image":
            try:
                import base64

                b64 = base64.b64encode(path.read_bytes()).decode()
                summary = self.llm.describe_image(b64, question or "Describe this image in detail, including any text.")
            except OSError as exc:
                summary = f"(couldn't open image: {exc})"
        else:
            summary = info.get("note", "I can see the file but not its contents yet.")
        info["summary"] = summary
        return info


# ---------------------------------------------------------------------------
# tools exposed to the model
# ---------------------------------------------------------------------------
def make_file_tools(local_search: LocalSearch, analyzer: FileAnalyzer) -> List[Tool]:
    def _search(args: str, ctx: ToolContext) -> ToolResult:
        if not ctx.permissions.check(Capability.READ_FILES, "to search your files", target=None):
            return ToolResult(False, "You haven't allowed me to read files.")
        hits = local_search.search(args, limit=12)
        if not hits:
            return ToolResult(True, f"No files matched '{args}' in the folders I'm allowed to see.")
        lines = [f"- {h['path']} ({h['size']})" + (f" — {h['snippet']}" if h["snippet"] else "") for h in hits]
        return ToolResult(True, "\n".join(lines), {"hits": hits})

    def _read(args: str, ctx: ToolContext) -> ToolResult:
        parts = args.split("|", 1)
        path = Path(parts[0].strip()).expanduser()
        question = parts[1].strip() if len(parts) > 1 else ""
        if not path.exists():
            return ToolResult(False, f"No such file: {path}")
        if not ctx.permissions.check(Capability.READ_FILES, "to read this file", target=str(path)):
            return ToolResult(False, f"You haven't allowed me to read {path.name}.")
        info = analyzer.analyze(path, question)
        text = info.get("text", "")
        out = f"{path.name} — {info['kind']} {json.dumps(info.get('meta', {}), default=str)}\n"
        if info.get("stats"):
            out += "Column stats: " + json.dumps(info["stats"], default=str)[:1500] + "\n"
        if info.get("summary"):
            out += "Analysis: " + info["summary"] + "\n"
        if text:
            out += "Excerpt:\n" + text[:3000]
        return ToolResult(True, out, {"info": {k: v for k, v in info.items() if k != "text"}})

    def _write(args: str, ctx: ToolContext) -> ToolResult:
        if "|" not in args:
            return ToolResult(False, "Usage: path | content")
        raw_path, content = args.split("|", 1)
        path = Path(raw_path.strip()).expanduser()
        if not ctx.permissions.check(Capability.WRITE_FILES, "to write this file", target=str(path)):
            return ToolResult(False, f"You haven't allowed me to write {path}.")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content.lstrip("\n"), encoding="utf-8")
        return ToolResult(True, f"Wrote {len(content)} characters to {path}")

    def _list(args: str, ctx: ToolContext) -> ToolResult:
        path = Path(args.strip() or ".").expanduser()
        if not path.is_dir():
            return ToolResult(False, f"Not a folder: {path}")
        if not ctx.permissions.check(Capability.READ_FILES, "to list this folder", target=str(path)):
            return ToolResult(False, f"You haven't allowed me to look inside {path}.")
        entries = []
        for child in sorted(path.iterdir())[:200]:
            if ctx.permissions.is_private(child):
                continue
            entries.append(("📁 " if child.is_dir() else "📄 ") + child.name)
        return ToolResult(True, "\n".join(entries) or "(empty)")

    return [
        Tool("search_files", "Search file names and contents in the folders you're allowed to see.", "what to look for", _search),
        Tool("read_file", "Read and analyse a file (text, code, PDF, DOCX, CSV, JSON, images with a vision model).", "path | optional question", _read),
        Tool("write_file", "Create or overwrite a file (asks permission outside Vidit's folder).", "path | content", _write, dangerous=True),
        Tool("list_folder", "List what's inside a folder.", "folder path", _list),
    ]
